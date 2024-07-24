import io
from openpyxl import load_workbook
from flask import current_app
from typing import Dict, Any, List, Optional
from openpyxl.worksheet.worksheet import Worksheet
from werkzeug.datastructures import FileStorage
from datetime import datetime
from dateutil import parser

from model.Case import CaseState

def func_header() -> Dict[str, str]:
    return {
            '用例编号': 'case_id',
            '用例名称': 'case_name',
            '测试环境': 'case_env',
            '所属模块': 'case_module',
            '前置条件': 'case_preconditions',
            '操作步骤': 'case_steps',
            '预期结果': 'case_expected_outcome',
            '测试结果': 'case_actual_outcome',
            '创建时间': 'case_create_time',
        }

'''
对于excel的处理本质上就是处理一个二维数组
整张表是一个数组
每一行也是一个数组
每一列是数组里面的值
合并单元格的本质就是二维数组当中的某个索引的数组里面存在空值.这个值和上一个数组的值相同
合并的单元格的值存储位置是左上角
'''
# 解析合并单元格 -> 声明一个字典.将单元格坐标作为key,
def parse_merged_cells(sheet: Worksheet) -> Dict[str, Any]:
    merge_cell_coordinate_value = {}

    for merge_range in sheet.merged_cells.ranges:
        top_left_value: str = str(
            sheet.cell(merge_range.min_row, merge_range.min_col).value
            ) if sheet.cell(merge_range.min_row, merge_range.min_col).value is not None else ''
        
        # 取出单元格坐标作为key,左上角值作为value
        for row in range(merge_range.min_row, merge_range.max_row + 1):
            for col in range(merge_range.min_col, merge_range.max_col + 1):
                merge_cell_coordinate_value[sheet.cell(row, col).coordinate] = top_left_value

    return merge_cell_coordinate_value

# datetime类型处理上转为字符串让json可以序列化
# 获取所有单元格数据.返回一个数组
def get_all_cell_data(sheet: Worksheet, excel_header: List[str], change_hander: Dict[str, str]) -> List[Dict[str, Any]]:
    all_cell_data = []
    merge_cells_values = {}

    merge_cells_values = parse_merged_cells(sheet)
    for row in sheet.iter_rows(min_row=2, max_row=sheet.max_row):
        each_row_data = {}
        # 每一行都和表头组合起来
        for cell, header in zip(row, excel_header):
            cell_value: Optional[str] = merge_cells_values.get(cell.coordinate, cell.value) # type: ignore
            english_header: str = change_hander.get(header, header)

            if english_header not in each_row_data:
                    each_row_data[english_header] = cell_value if cell_value is not None else ""
            else:
                each_row_data[english_header] = cell_value if cell_value is not None else each_row_data[english_header]

        if CaseState.from_str_get_state(each_row_data["case_actual_outcome"]) == CaseState.UNKNOWN.value:
            each_row_data["case_actual_outcome"] = CaseState.WAITING.value
        elif each_row_data["case_actual_outcome"] == '':
            each_row_data["case_actual_outcome"] = CaseState.WAITING.value

        if isinstance(each_row_data["case_create_time"], datetime):
            each_row_data['case_create_time'] = each_row_data['case_create_time'].isoformat()
        else:
            dt: datetime = parser.parse(each_row_data["case_create_time"])
            each_row_data["case_create_time"] = dt.strftime("%Y-%m-%d")
        all_cell_data.append(each_row_data)
    return all_cell_data

def parse_case_template_excel(file: FileStorage, sheet_name: str = "") -> List[Dict[str, Any]]:
    try:
        file_stream = io.BytesIO(file.read())
        workbook = load_workbook(file_stream, data_only=True)
        sheet = workbook[sheet_name] if sheet_name else workbook.active
        headers = [
            cell.value
            for cell in next(sheet.iter_rows(min_row=1, max_row=1)) # type: ignore
        ]
        header_dict = func_header()
        data = get_all_cell_data(sheet, headers, header_dict) # type: ignore
        file_stream.close()
        return data
    except Exception as e:
        current_app.logger.error(f"fail to parse case template: {str(e)}")
        raise IOError(f"Unable to parse Excel file") from e

